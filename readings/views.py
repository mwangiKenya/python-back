from django.db import transaction
from django.http import JsonResponse, HttpResponse
from api.models import readings,Billings
import json
from django.views.decorators.csrf import csrf_exempt
from datetime import date
import os
from django.conf import settings
from openpyxl import load_workbook
import pandas as pd
from django.db import transaction
from django.db.models import F
from django.db.models.functions import Coalesce
from django.utils import timezone
from django.views.decorators.http import require_POST


def read_data(request):
    data = list(readings.objects.values())
    return JsonResponse(data, safe=False)

@csrf_exempt
def submit_new_reading(request):
    if request.method != "POST":
        return JsonResponse({"error": "Invalid request"}, status=400)
    try:
        data = json.loads(request.body)
        updates = data if isinstance(data, list) else [data]
        with transaction.atomic():
            for item in updates:
                reading = readings.objects.select_for_update().get(user_id=item["user_id"])
                cur_user = item.get("cur_user")
                cur_sup = item.get("cur_sup") 
                if cur_user is not None:
                    cur_user = int(cur_user)
                    reading.cur_user = cur_user
                    reading.units_used = max(0, cur_user - (reading.prev_user or 0))
                if cur_sup is not None:
                    reading.cur_sup = cur_sup
                reading.mid_user = item.get("mid_user", reading.mid_user)
                reading.mid_sup = item.get("mid_sup", reading.mid_sup)
                reading.save()
                rate = reading.rate or 0
                units = reading.units_used or 0
                calculated_bill = units * rate
                final_bill = max(300.0, float(calculated_bill))
                billing_record, created = Billings.objects.select_for_update().get_or_create(
                    user_id=reading.user_id,
                    defaults={
                        "name": reading.name,
                        "phone": reading.phone,
                        "billed_on": date.today(),
                        "units_used": units,
                        "rate": rate,
                        "bill": final_bill,
                        "paid": 0.0,
                        "b_cd": 0.0,
                        "bal": final_bill,
                        "status": "PENDING",
                        "prev_user": reading.prev_user or 0,
                        "cur_user": reading.cur_user or 0,
                        "sms_name": reading.name[:50],
                        "grp": reading.grp,
                        "parent": reading.parent,
                        "gen_name": reading.gen_name
                    }
                )
                print(reading)
                if not created:
                    carried_forward = float(billing_record.bal or 0.0)
                    new_balance = final_bill + carried_forward
                    billing_record.name = reading.name
                    billing_record.phone = reading.phone
                    billing_record.billed_on = date.today()
                    billing_record.units_used = units
                    billing_record.rate = rate
                    billing_record.b_cd = carried_forward
                    billing_record.bill = final_bill
                    billing_record.paid = 0
                    billing_record.bal = new_balance
                    billing_record.prev_user = reading.prev_user or 0
                    billing_record.cur_user = reading.cur_user or 0
                    billing_record.grp = reading.grp
                    billing_record.parent = reading.parent
                    billing_record.gen_name = reading.gen_name
                    billing_record.save()
        return JsonResponse({"message": "Saved successfully"})
    except Exception as e:
        return JsonResponse({"error": str(e)}, status=500)

def download_readings_template(request):
    template_path = os.path.join(settings.BASE_DIR, "templates", "readings_template.xlsx")
    wb = load_workbook(template_path)
    ws = wb.active
    readings_data = readings.objects.all().values_list(
        "user_id", "name", "phone", "metre_num", "prev_user", "prev_sup"
    )
    row = 2
    for r in readings_data:
        for col_idx, value in enumerate(r, start=1):
            ws.cell(row=row, column=col_idx, value=value)
        for col_idx in range(7, 11):
            ws.cell(row=row, column=col_idx, value=None)
        row += 1
    while row <= ws.max_row:
        for col_idx in range(1, 11):
            ws.cell(row=row, column=col_idx, value=None)
        row += 1
    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="readings_template.xlsx"'
    wb.save(response)
    return response

@csrf_exempt
@require_POST
def upload_readings_excel(request):
    file = request.FILES.get("file")
    if not file:
        return JsonResponse({"error": "No file uploaded"}, status=400)
    try:
        df = pd.read_excel(file)
        if not {"user_id", "cur_user"} <= set(df.columns):
            return JsonResponse({"error": "Excel must contain user_id and cur_user columns"}, status=400)
        processed, errors = 0, []
        for i, row in df.iterrows():
            if pd.isna(row["user_id"]) or pd.isna(row["cur_user"]):
                continue
            try:
                with transaction.atomic():
                    reading = readings.objects.select_for_update().get(user_id=int(row["user_id"]))
                    cur_user = int(row["cur_user"])
                    cur_sup = int(row["cur_sup"]) if "cur_sup" in df and pd.notna(row["cur_sup"]) else None
                    if cur_user < (reading.prev_user or 0) or (cur_sup is not None and cur_sup < (reading.prev_sup or 0)):
                        raise ValueError("current reading is less than previous")
                    reading.cur_user = cur_user
                    reading.units_used = max(0, cur_user - (reading.prev_user or 0))
                    if cur_sup is not None:
                        reading.cur_sup = cur_sup
                    reading.save()
                    rate, units = reading.rate or 0, reading.units_used or 0
                    final_bill = max(300.0, float(units * rate))
                    vals = dict(name=reading.name, phone=reading.phone, billed_on=date.today(), units_used=units, rate=rate,
                                bill=final_bill, prev_user=reading.prev_user or 0, cur_user=reading.cur_user or 0,
                                grp=reading.grp, parent=reading.parent)
                    b, created = Billings.objects.select_for_update().get_or_create(
                        user_id=reading.user_id,
                        defaults={**vals, "paid": 0.0, "b_cd": 0.0, "bal": final_bill, "status": "PENDING", "sms_name": reading.name[:50]})
                    if not created:
                        carried = float(b.bal or 0.0)
                        for k, v in vals.items():
                            setattr(b, k, v)
                        b.b_cd, b.paid, b.bal = carried, 0, final_bill + carried
                        b.save()
                    processed += 1
            except Exception as e:
                errors.append(f"Row {i}: {e}")
        return JsonResponse({"message": f"{processed} readings billed, {len(errors)} failed", "processed_rows": processed, "errors": errors[:10]})
    except Exception as e:
        return JsonResponse({"error": str(e)}, status=500)

@csrf_exempt
@require_POST
def shift_readings(request):
    with transaction.atomic():
        n = readings.objects.update(
            prev_user=Coalesce('cur_user', 'prev_user'),
            prev_sup=Coalesce('cur_sup', 'prev_sup'),
            prev_date=Coalesce('cur_date', 'prev_date'),
            cur_user=None,
            cur_sup=None,
            cur_date=timezone.localdate(),
            version=F('version') + 1,
        )
    return JsonResponse({"message": f"{n} readings shifted"})